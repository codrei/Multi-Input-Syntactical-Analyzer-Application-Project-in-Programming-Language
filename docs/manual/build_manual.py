"""Build the User Manual and System Documentation as a Word file.

    pip install python-docx
    python docs/manual/build_manual.py

Writes docs/SyntaxLens_User_Manual.docx.  The PDF version is exported from
that file with Microsoft Word (File > Save As > PDF), so both match exactly.
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
OUTPUT = ROOT / "docs" / "SyntaxLens_User_Manual.docx"
SCREENSHOT = ROOT / "docs" / "images" / "web-app.png"

LIVE_URL = "https://multi-input-syntactical-analyzer-ap.vercel.app"
REPO_URL = "https://github.com/codrei/Multi-Input-Syntactical-Analyzer-Application-Project-in-Programming-Language"
MEMBERS = "Marco Andrei R. Belen · Gabriel M. Caporado · Ariz Luiz L. Guillarte · Brix Martin M. Paña"

ACCENT = RGBColor(0x2B, 0x2F, 0x8F)
MUTED = RGBColor(0x55, 0x5B, 0x6E)
HEADER_FILL = "E8EAF6"
CODE_FILL = "F4F5F8"


# ----------------------------------------------------------------- helpers

def shade(element, fill: str) -> None:
    """Give a table cell or paragraph a background color."""
    properties = element.get_or_add_tcPr() if element.tag.endswith("}tc") else element.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def rich(paragraph, text: str, size: float | None = None) -> None:
    """Add text where **bold**, *italic* and `code` are marked, like in Markdown."""
    for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*\s][^*]*\*)", text):
        if not part:
            continue
        if part.startswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        elif part.startswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt((size or 10.5) - 1)
        else:
            run = paragraph.add_run(part)
        if size and not part.startswith("`"):
            run.font.size = Pt(size)


def para(doc, text: str, size: float | None = None, after: float = 4, style: str | None = None):
    paragraph = doc.add_paragraph(style=style)
    rich(paragraph, text, size)
    paragraph.paragraph_format.space_after = Pt(after)
    return paragraph


def bullets(doc, items: list[str], style: str = "List Bullet") -> None:
    for item in items:
        paragraph = doc.add_paragraph(style=style)
        rich(paragraph, item)
        paragraph.paragraph_format.space_after = Pt(1.5)


def heading(doc, text: str, level: int = 1) -> None:
    paragraph = doc.add_heading(text, level=level)
    paragraph.paragraph_format.space_before = Pt(10 if level == 1 else 6)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.keep_with_next = True


def code(doc, lines: list[str]) -> None:
    """A shaded block of monospaced text, one paragraph per line."""
    for number, line in enumerate(lines):
        paragraph = doc.add_paragraph()
        run = paragraph.add_run(line if line else " ")
        run.font.name = "Consolas"
        run.font.size = Pt(8.5)
        shade(paragraph._p, CODE_FILL)
        fmt = paragraph.paragraph_format
        fmt.space_after = Pt(0)
        fmt.space_before = Pt(0)
        fmt.left_indent = Inches(0.15)
        fmt.keep_with_next = number < len(lines) - 1
    doc.paragraphs[-1].paragraph_format.space_after = Pt(6)


def table(doc, rows: list[list[str]], widths: list[float], size: float = 9.5) -> None:
    """A table with a shaded header row; widths in inches."""
    grid = doc.add_table(rows=len(rows), cols=len(widths))
    grid.style = "Table Grid"
    grid.alignment = WD_TABLE_ALIGNMENT.CENTER
    grid.autofit = False
    for r, row in enumerate(rows):
        # A row never breaks across two pages.
        grid.rows[r]._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for c, text in enumerate(row):
            cell = grid.cell(r, c)
            cell.width = Inches(widths[c])
            cell.paragraphs[0].text = ""
            rich(cell.paragraphs[0], f"**{text}**" if r == 0 else text, size)
            cell.paragraphs[0].paragraph_format.space_after = Pt(0)
            if r == 0:
                shade(cell._tc, HEADER_FILL)
    for c, width in enumerate(widths):
        grid.columns[c].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def rule(doc) -> None:
    """A thin line under the title block (a paragraph border, not a table)."""
    paragraph = doc.add_paragraph()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for key, value in (("val", "single"), ("sz", "8"), ("space", "1"), ("color", "2B2F8F")):
        bottom.set(qn(f"w:{key}"), value)
    borders.append(bottom)
    paragraph._p.get_or_add_pPr().append(borders)
    paragraph.paragraph_format.space_after = Pt(6)


# --------------------------------------------------------------- document

def build() -> Path:
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    for side in ("left_margin", "right_margin"):
        setattr(section, side, Inches(0.85))
    section.top_margin = section.bottom_margin = Inches(0.75)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    for name, size in (("Heading 1", 14), ("Heading 2", 11.5), ("Title", 26)):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = ACCENT
        style.font.bold = True

    footer = section.footer.paragraphs[0]
    footer.text = "SyntaxLens · User Manual and System Documentation · COM 241"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.runs[0].font.size = Pt(8.5)
    footer.runs[0].font.color.rgb = MUTED

    # ---- title block
    title = doc.add_paragraph(style="Title")
    title.add_run("SyntaxLens")
    title.paragraph_format.space_after = Pt(0)
    subtitle = para(doc, "**Multi-Input Syntactical Analyzer: User Manual and System Documentation**", 13, 6)
    subtitle.runs[0].font.color.rgb = ACCENT
    para(doc, f"**COM 241 Final Project** · October 2026 · Group members: {MEMBERS}", 10, 2)
    para(doc, f"Live app: {LIVE_URL}", 10, 1)
    para(doc, f"Source code: {REPO_URL}", 10, 2)
    rule(doc)

    # ---- 1
    heading(doc, "1. Overview")
    para(doc, "SyntaxLens checks source code for syntax errors **without running it**. It splits the "
              "code into tokens, checks them against syntax rules, and prints a report that either "
              "confirms the code is valid or lists each error with its line, column, category, "
              "explanation and a suggested fix.")
    bullets(doc, [
        "**Three input modes:** a single line, a block of code, or an uploaded file.",
        "**Five languages:** Python, Java, C, C++ and C#, detected automatically (or chosen by the user).",
        "**Eight syntax checks:** the six in the specification (brackets, string literals, statement "
        "terminators, operators, control structure headers, identifier names) plus indentation & block "
        "structure, and illegal characters & malformed literals.",
        "**Two interfaces, one engine:** a web app (online, or offline on your own computer) and a "
        "command line. Both call the same analysis function, so they always agree.",
        "**Reports:** the specification's exact format (classic), or a detailed report with error "
        "codes, a pointer to the exact column, fix suggestions and totals for every category; "
        "also as JSON data.",
    ])

    # ---- 2
    heading(doc, "2. Getting Started")
    table(doc, [
        ["Way to run", "How", "What you need"],
        ["Online", f"Open {LIVE_URL.removeprefix('https://')}", "A browser and internet"],
        ["Windows, offline", "Double-click `run.bat` in the project folder; the app opens in the browser",
         "Python 3.10 or newer"],
        ["Command line", "`py -m syntaxlens` (menu) or `py -m syntaxlens file code.py`",
         "Python 3.10 or newer"],
    ], [1.25, 3.55, 1.95])
    para(doc, "Python is free from python.org (tick \"Add python.exe to PATH\" while installing). "
              "SyntaxLens uses only Python's standard library, so nothing else has to be installed. "
              "On macOS or Linux, write `python3` instead of `py`.")

    # ---- 3
    heading(doc, "3. Using the Three Input Modes")
    heading(doc, "3.1 In the web app", 2)
    bullets(doc, [
        "**Language:** leave it on *Auto-detect*, or choose the language in the menu at the top.",
        "**Mode 1, Single Line:** type one statement, e.g. `x = 10` or `System.out.println(\"Hello\");`, "
        "and press Enter. The example buttons fill in ready-made lines.",
        "**Mode 2, Code Block:** type or paste several lines into the editor and press *Analyze* "
        "(or Ctrl+Enter). The *Load a sample* menu offers the specification's examples and one demo "
        "file per language. *Analyze as I type* re-checks the code while you edit.",
        "**Mode 3, Upload File:** drag a file onto the box, or click it to choose one (.py, .java, .c, "
        ".cpp, .h, .cs, .txt; up to 2 MB). Empty, binary (e.g. images) and oversized files, and files "
        "in other languages (e.g. .js), are refused with a message.",
    ], style="List Number")
    para(doc, "The results show a PASSED / FAILED banner, counts, and bar charts of errors per check "
              "and tokens per category. Errors are marked in the code with a red line, a dot in the "
              "margin and a wavy underline. The tabs give the **Errors** (click one to jump to its "
              "line), the **Line Breakdown**, all **Tokens** (filter by category) and the text "
              "**Report** (detailed or classic) with Copy and Download (.txt, .json).")
    doc.add_picture(str(SCREENSHOT), width=Inches(5.9))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption = para(doc, "Figure 1. The specification's Example 2 in the web app: three errors, "
                        "each marked in the code and explained on the right.", 9, 6)
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.runs[0].font.color.rgb = MUTED

    heading(doc, "3.2 On the command line", 2)
    table(doc, [
        ["Input mode", "Command"],
        ["Menu (all modes)", "`py -m syntaxlens`, then choose 1, 2 or 3"],
        ["1. Single line", "`py -m syntaxlens line \"x = 10\"`"],
        ["2. Code block", "`py -m syntaxlens block`, paste the code, end with a line holding only `.`"],
        ["3. File", "`py -m syntaxlens file samples\\spec_example_2.py` (several files are allowed)"],
    ], [1.6, 5.15])
    para(doc, "Options: `--lang python|java|c|cpp|csharp`, `--style classic` (the specification's "
              "exact layout), `--format json`, `--output report.txt`, `--no-color`. The exit code is "
              "0 for valid code, 1 when syntax errors are found and 2 when the input cannot be read.")
    heading(doc, "3.3 Reading an error", 2)
    para(doc, "Each error names its line, category and cause. The detailed style adds the error code, "
              "the check that found it, a caret under the column and a fix:")
    code(doc, [
        "[ERROR 3] Line 4: y = 20 + * 5",
        "  - Category: Invalid Operator Sequence",
        "  - Details: Consecutive binary operators '+ *' without an operand in between.",
        "  - Location: line 4, column 10 (check 4: Operator Syntax, code E401)",
        "      4 | y = 20 + * 5",
        "        |          ^",
        "  - Fix: Remove one of the operators, or put a value between them.",
    ])

    # ---- 4
    heading(doc, "4. System Architecture")
    para(doc, "SyntaxLens is a pipeline of seven small stages, each in its own Python module. "
              "The output of one stage is the input of the next:")
    table(doc, [
        ["Stage", "Module", "What it does"],
        ["1. Source", "`source.py`", "Decodes the input (UTF-8, Windows-1252, byte-order marks), turns "
                                     "Windows line endings into one form, refuses blank, binary and "
                                     "oversized input."],
        ["2. Detect", "`detect.py`", "Chooses the language: the user's choice, else the file extension, "
                                     "else weighted clues in the code (e.g. `System.out.println` for Java)."],
        ["3. Lexer", "`lexer.py`", "Splits the text into tokens with one regular expression and sorts "
                                   "them into 8 categories; marks lexical problems on the tokens."],
        ["4. Statements", "`structure.py`", "Groups tokens into statements (Python: logical lines and "
                                            "indentation; C family: `;` `{` `}`) and labels each one."],
        ["5. Checks", "`checks/`", "Eight independent modules. Each error type belongs to exactly one "
                                   "check, so nothing is reported twice."],
        ["6. Filter", "`analyzer.py`", "Keeps the root cause of each mistake and drops the follow-on "
                                       "errors it causes in the same statement."],
        ["7. Report", "`report/`, `cli.py`, `web/`", "Text report (classic or detailed), JSON, the "
                                                     "command line and the web app."],
    ], [1.05, 1.55, 4.15])
    para(doc, "**Languages are data, not code.** Each language is a *profile* (`profiles/`) with its "
              "spelling (keywords, operators, comments, number and string formats) and its "
              "`LanguageRules`: switches for what it allows, such as pointers in C or properties in "
              "C#. The shared code never asks which language it is analyzing, and a test enforces "
              "this. The web app is a small WSGI application from Python's standard library; it "
              "calls the same analysis function as the command line and runs both locally and on "
              "Vercel.")

    # ---- 5
    heading(doc, "5. Parsing Logic")
    heading(doc, "5.1 Tokenizing", 2)
    para(doc, "All token patterns of a language are joined into **one master regular expression**, "
              "one named group per token type. Starting at the first character, the lexer asks which "
              "group matches, records that token and continues after it: a single left-to-right pass. "
              "Three rules make it reliable:")
    bullets(doc, [
        "**Order matters:** comments and strings are tried before operators (`//` is a comment, "
        "not two slashes) and operators longest first (`**=` before `**` before `*`).",
        "**Strings and comments are single tokens,** so code-like text inside them, such as "
        "`\"if (x > 0)\"`, is never checked as code.",
        "**It never stops:** a character the language does not allow becomes an *invalid* token, "
        "and an unclosed string ends at the end of its line.",
    ])
    para(doc, "Every token falls into one of 8 categories, which the report lists with totals: "
              "keywords, identifiers, operators, numbers, strings and characters, delimiters, "
              "comments, and invalid tokens. The specification's Example 1 gives exactly 14 tokens.")
    heading(doc, "5.2 The eight checks", 2)
    table(doc, [
        ["#", "Check", "How it works", "Example error"],
        ["1", "Delimiters & brackets", "A stack: each opening bracket is pushed; a closing bracket must "
                                       "match the one on top. Leftovers were never closed.", "`if (x > 5`"],
        ["2", "String & char literals", "The lexer marks literals with no closing quote; a char "
                                        "literal must hold one character.", "`print(\"hi)`"],
        ["3", "Statement terminators", "Missing `;` in Java, C, C++ and C#; missing `:` after a "
                                       "Python header.", "`int x = 10`"],
        ["4", "Operator syntax", "A two-state machine that expects a value or an operator; `+5` is "
                                 "a sign, `+ *` is an error.", "`y = 20 + * 5`"],
        ["5", "Control headers", "A small grammar per header: if, for, while, switch, def, class.",
         "`for i range(10):`"],
        ["6", "Identifier naming", "Names start with a letter or `_` and are not reserved words.",
         "`int 2nd = 3;`"],
        ["7", "Indentation & blocks", "A stack of indentation levels (as in CPython); `else`, `case`, "
                                      "`break` must be inside the right block.", "unexpected indent"],
        ["8", "Illegal characters & literals", "Illegal characters, malformed numbers, unclosed "
                                               "comments, curly quotes.", "`3.14.15`"],
    ], [0.3, 1.5, 3.4, 1.55])
    heading(doc, "5.3 Error recovery", 2)
    para(doc, "The analyzer does not stop at the first error. (1) The lexer never stops. (2) At each "
              "line break inside an unfinished statement it asks whether the line ends with a complete "
              "value while the next line starts something new; if so, the statement ends there and the "
              "missing `)` or `;` is reported on this line. (3) A filter keeps only the root cause: an "
              "unclosed string hides the bracket and operator errors it caused in the same statement. "
              "This is why the specification's Example 2 gives exactly its three errors. Python's own "
              "parser reports only the first one, and a naive checker reports about six.")

    # ---- 6
    heading(doc, "6. Regular Expressions and Libraries")
    table(doc, [
        ["Token", "Pattern (simplified)"],
        ["Identifier (Python, C, C++)", "`[^\\W\\d]\\w*`  (a letter or _, then letters, digits or _)"],
        ["Identifier (Java)", "`(?:[^\\W\\d]|\\$)(?:\\w|\\$)*`  (also allows $)"],
        ["Number (Python)",
         "`0[xX](?:_?[0-9a-fA-F])+` | `[1-9](?:_?[0-9])*` | floats, from CPython's tokenizer"],
        ["String", "`\"(?:[^\"\\\\\\n]|\\\\.)*\"`  unclosed: the same without the closing quote"],
        ["Comments", "`#[^\\n]*`   `//[^\\n]*`   `/\\*[\\s\\S]*?\\*/`"],
        ["Operators", "one alternation, longest first: `\\*\\*=|//=|==|\\+|-|...`"],
        ["Anything else", "`.`  (one illegal character becomes an invalid token)"],
    ], [1.9, 4.85])
    para(doc, "SyntaxLens uses only Python's standard library (`re`, `dataclasses`, `enum`, `wsgiref`). "
              "No parser generator or AST library is used: they stop at the first error, so they "
              "cannot produce the specification's Example 2. Python's `tokenize` module is used only "
              "in the tests, to confirm that our Python lexer produces the same tokens. The web "
              "editor is CodeMirror 5 (MIT license), stored in the repository so the app works offline.")

    # ---- 7
    heading(doc, "7. Testing")
    bullets(doc, [
        "**610 automated tests** run on Windows and Linux with Python 3.10 to 3.14 on every change.",
        "Both example reports of the specification are reproduced line for line; only the "
        "*Target Syntax Rule* line differs, because it names the detected language.",
        "**No false alarms:** 19 realistic valid programs in the five languages must give zero "
        "findings; the Python, Java, C and C++ ones are also checked by the real compilers "
        "(CPython, javac, gcc, g++).",
        "**Catch rate:** `tools/measure_accuracy.py` plants 1,601 known mistakes, one at a time, into "
        "those programs (a missing `)` or `;`, an unclosed string, `+ /`, a name starting with a digit, "
        "...). **98.3%** are reported on the right line. A test keeps the rate above 97%.",
        "Every error type has a test in every language it applies to.",
        "Thousands of random inputs never crash the analyzer; 10,000 lines take under half a second.",
    ])

    # ---- 8
    heading(doc, "8. Error Code Reference")
    table(doc, [
        ["Codes", "Check", "Meaning"],
        ["E101–E103", "1 Delimiters", "unclosed, unexpected or mismatched bracket"],
        ["E201–E204", "2 Literals", "unclosed string or char literal; empty or multi-character char literal"],
        ["E301–E302", "3 Terminators", "missing semicolon `;`, or missing colon `:` after a Python header"],
        ["E401–E408, W406", "4 Operators", "operators in a row, missing operand or operator, invalid "
                                           "assignment target, value computed but not used, empty argument"],
        ["E501–E510, W505, W509", "5 Headers", "missing condition or parentheses, bad for/foreach "
                                               "header, else with a condition, bad function header; "
                                               "`=` in a condition, empty statement after a header"],
        ["E601–E604", "6 Identifiers", "name starts with a digit, contains `-` or a space, is a "
                                       "reserved word, or is missing"],
        ["E701–E708", "7 Blocks", "unexpected or missing indentation, mixed tabs and spaces, `else` "
                                  "without `if`, `break` outside a loop, `case` outside `switch`"],
        ["E801–E805", "8 Lexical", "illegal character, malformed number, unclosed comment, curly "
                                   "quotes, invisible look-alike space"],
    ], [1.55, 1.2, 4.0], size=9)
    para(doc, "Codes starting with **W** are warnings: valid code that is probably a mistake. They "
              "do not make the analysis fail.", 9.5)

    # ---- 9
    heading(doc, "9. Troubleshooting")
    table(doc, [
        ["Problem", "What to do"],
        ["`py` is not recognized", "Install Python 3.10 or newer from python.org and tick \"Add python.exe "
                                   "to PATH\". `run.bat` also tries `python` if `py` is missing."],
        ["The browser does not open", "Look at the console window: it prints the address, e.g. "
                                      "`http://127.0.0.1:8000`. Open it by hand. Keep the window open "
                                      "while you use the app; Ctrl+C stops it."],
        ["Port 8000 is already in use", "Nothing to do: the next free port (8001, 8002, ...) is used "
                                        "and printed. Or choose one: `py -m syntaxlens web --port 9000`."],
        ["The wrong language is detected", "Very short code can look like several languages. Choose the "
                                           "language in the menu, or add `--lang java` on the command line."],
        ["\"... is empty\" or \"looks like a binary file\"",
         "Upload the source file itself (.py, .java, .c, .cpp, .cs or .txt), "
         "not a Word document, PDF or image."],
        ["The online app is slow the first time", "The server sleeps when nobody uses it; the first "
                                                  "analysis wakes it up in a few seconds. The offline "
                                                  "app has no delay."],
    ], [2.1, 4.65], size=9.5)

    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(f"Wrote {build()}")
